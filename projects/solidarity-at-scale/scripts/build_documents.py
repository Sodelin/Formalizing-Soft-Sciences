"""Render the project Markdown to editable DOCX. Requires python-docx.

PDF rendering is a separate LibreOffice step; see HANDOFF.md.
Only the restricted Markdown constructs used in these three documents are parsed.
"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

PROJECT = Path(__file__).resolve().parents[1]

def hyperlink(paragraph, label, url):
    rel = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement('w:hyperlink'); link.set(qn('r:id'), rel)
    run = OxmlElement('w:r'); props = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '000000'); props.append(color)
    underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single'); props.append(underline)
    run.append(props); text = OxmlElement('w:t'); text.text = label; run.append(text)
    link.append(run); paragraph._p.append(link)

def inline(paragraph, text):
    tokens = re.split(r'(\*\*.*?\*\*|`[^`]+`|\[[^\]]+\]\(https?://[^)]+\))', text)
    for token in tokens:
        if not token: continue
        if token.startswith('**') and token.endswith('**'):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith('`') and token.endswith('`'):
            run=paragraph.add_run(token[1:-1]);run.font.name='Courier New';run.font.size=Pt(10)
        elif re.match(r'\[[^\]]+\]\(https?://',token):
            label,url=re.match(r'\[([^\]]+)\]\((.*)\)',token).groups();hyperlink(paragraph,label,url)
        else: paragraph.add_run(token)

def setup(doc, short_title):
    sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=Inches(.68);sec.bottom_margin=Inches(.68)
    sec.left_margin=Inches(.8);sec.right_margin=Inches(.8)
    sec.header_distance=Inches(.3);sec.footer_distance=Inches(.3)
    styles=doc.styles
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        styles[name].font.name='Times New Roman';styles[name].font.color.rgb=RGBColor(0,0,0)
        el=styles[name].element
        for border in el.xpath('./w:pPr/w:pBdr'):border.getparent().remove(border)
        for fonts in el.xpath('./w:rPr/w:rFonts'):
            for key in list(fonts.attrib):
                if 'theme' in key.lower():del fonts.attrib[key]
    styles['Normal'].font.size=Pt(11.5)
    styles['Normal'].paragraph_format.line_spacing=1.06
    styles['Normal'].paragraph_format.space_after=Pt(6)
    styles['Normal'].paragraph_format.widow_control=True
    styles['Title'].font.size=Pt(20);styles['Title'].font.bold=True
    styles['Title'].paragraph_format.space_after=Pt(5)
    styles['Subtitle'].font.size=Pt(12.5);styles['Subtitle'].paragraph_format.space_after=Pt(8)
    for name,size in [('Heading 1',13),('Heading 2',11.5),('Heading 3',11)]:
        styles[name].font.size=Pt(size);styles[name].font.bold=True
        styles[name].paragraph_format.space_before=Pt(11)
        styles[name].paragraph_format.space_after=Pt(5)
        styles[name].paragraph_format.keep_with_next=True
    header=sec.header.paragraphs[0];header.text='SOLIDARITY AT SCALE  |  '+short_title.upper()
    header.style='Normal';header.runs[0].font.size=Pt(8)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Initial research version 0.1  •  ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    for run in footer.runs:run.font.size=Pt(8)
    doc.core_properties.author='Codex research assistance for Nolan Downard'
    doc.core_properties.subject='Identity, collective action, political institutions, and scoped Lean models'

def table(doc, lines):
    cells=[[x.strip() for x in line.strip().strip('|').split('|')] for line in lines]
    cells=[row for row in cells if not all(re.fullmatch(r':?-+:?',c) for c in row)]
    t=doc.add_table(rows=1,cols=len(cells[0])); t.autofit=False
    widths={2:[2.2,4.7],3:[1.7,2.6,2.6],4:[1.55,1.65,1.65,2.05]}.get(len(cells[0]),[6.9/len(cells[0])]*len(cells[0]))
    if cells[0][1]=='Score':widths=[2.3,.6,4.0]
    for col,width in zip(t.columns,widths):col.width=Inches(width)
    for n,row in enumerate(cells):
        dest=t.rows[0] if n==0 else t.add_row()
        if n==0:
            repeat=OxmlElement('w:tblHeader');dest._tr.get_or_add_trPr().append(repeat)
        no_split=OxmlElement('w:cantSplit');dest._tr.get_or_add_trPr().append(no_split)
        for cell,txt,width in zip(dest.cells,row,widths):
            cell.width=Inches(width);p=cell.paragraphs[0];inline(p,txt)
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if txt.isdigit():p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
            p.paragraph_format.line_spacing=1.0
            p.paragraph_format.keep_with_next=n<len(cells)-1
            for run in p.runs:run.font.size=Pt(9.5);run.bold=n==0
            props=cell._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'75');e.set(qn('w:type'),'dxa');margins.append(e)
            props.append(margins)
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            props.append(borders)
            if n==0:
                shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'F2F2F2');props.append(shading)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

def body(doc, markdown, report=False, references=False):
    lines=markdown.splitlines();i=0;seen_title=False;seen_subtitle=False
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):rows.append(lines[i]);i+=1
            table(doc,rows);continue
        if line.startswith('# '):
            if references:doc.add_heading('References and access notes',level=1)
            else:doc.add_paragraph(line[2:],style='Title');seen_title=True
        elif line.startswith('## '):
            txt=line[3:]
            if seen_title and not seen_subtitle and not references:
                doc.add_paragraph(txt,style='Subtitle');seen_subtitle=True
            else:
                p=doc.add_heading(txt,level=1)
                if report and txt=='1 Abstract':p.paragraph_format.page_break_before=True
                if references:
                    p.paragraph_format.space_before=Pt(8);p.paragraph_format.space_after=Pt(3)
                    p.runs[0].font.size=Pt(10.5)
        elif line.startswith('### '):doc.add_heading(line[4:],level=2)
        else:
            parts=[line]
            while i+1<len(lines) and lines[i+1].strip() and not re.match(r'^(#|\||\d+\. |[-*] )',lines[i+1].strip()):
                i+=1;parts.append(lines[i].strip())
            txt=' '.join(parts)
            if references:
                txt=re.sub(r'URL: (https?://\S+)',r'[Source link](\1)',txt)
                txt=re.sub(r'DOI: (https?://doi.org/\S+)',r'[DOI](\1)',txt)
            p=doc.add_paragraph();inline(p,txt)
            if references:
                p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.0
                p.paragraph_format.keep_with_next=not txt.startswith('Access:')
                for run in p.runs:run.font.size=Pt(10)
        i+=1

def create(source,stem,title,report=False,references='sources/references.md'):
    doc=Document();setup(doc,title)
    doc.core_properties.title='Solidarity at scale: '+title
    body(doc,(PROJECT/source).read_text(),report)
    doc.add_page_break()
    body(doc,(PROJECT/references).read_text(),references=True)
    output=PROJECT/'outputs'/f'{stem}.docx';output.parent.mkdir(exist_ok=True)
    doc.save(output);print(output)

if __name__=='__main__':
    create('paper/manuscript.md','solidarity-at-scale-paper','Research manuscript')
    create('report/research-report.md','solidarity-at-scale-report','Evidence and political decision report',True)
    create('formal/what-lean-proves.md','what-lean-proves','Lean reader guide',references='sources/formalization-prior-work.md')
